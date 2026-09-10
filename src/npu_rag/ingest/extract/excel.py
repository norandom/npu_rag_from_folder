"""Excel extractor: blank-row blocks with merged-cell attribution (task 4.4).

A row is empty only if every cell is empty and no valued merged range covers
it. Hidden sheets become ``Omission(HIDDEN_SHEET)`` and are never raised.
Every block carries a ``SheetLocator``. This module never imports vision,
state, or httpx.

``SourceFile.path`` is opened as given, including the Windows ``\\\\?\\`` form,
by reading bytes and handing them to openpyxl.
"""

from __future__ import annotations

import zipfile
from collections.abc import Sequence
from io import BytesIO

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.cell.cell import MergedCell  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]

from npu_rag.ingest.config import IngestConfig
from npu_rag.ingest.errors import ExtractionError
from npu_rag.ingest.extract.base import normalise
from npu_rag.ingest.types import (
    BlockSegment,
    Extracted,
    Omission,
    OmissionCategory,
    Segment,
    SheetLocator,
    SourceFile,
)

__all__ = ["ExcelExtractor"]


class ExcelExtractor:
    """Segment visible worksheets into labelled blocks. Conforms to ``Extractor``."""

    def extract(self, source: SourceFile, config: IngestConfig) -> Extracted:
        del config
        try:
            payload = source.path.read_bytes()
        except OSError as exc:
            raise ExtractionError(
                f"could not read {source.relative_path.as_posix()} as a workbook",
                stage="extraction",
                path=source.path,
            ) from exc
        try:
            workbook = load_workbook(BytesIO(payload), data_only=False)
        except (InvalidFileException, zipfile.BadZipFile, OSError, KeyError, ValueError) as exc:
            raise ExtractionError(
                f"could not read {source.relative_path.as_posix()} as a workbook",
                stage="extraction",
                path=source.path,
            ) from exc
        try:
            segments: list[Segment] = []
            omissions: list[Omission] = []
            for sheet in workbook.worksheets:
                if sheet.sheet_state != "visible":
                    omissions.append(
                        Omission(
                            category=OmissionCategory.HIDDEN_SHEET,
                            path=source.path,
                            reason=f"hidden sheet {sheet.title!r} skipped",
                        )
                    )
                    continue
                segments.extend(_blocks_for_sheet(sheet))
        finally:
            workbook.close()
        return Extracted(title=None, segments=tuple(segments), omissions=tuple(omissions))


def _row_is_empty(
    row_index: int,
    values: Sequence[object],
    valued_merged_ranges: Sequence[tuple[int, int]],
) -> bool:
    """A row is empty iff every cell is empty and no valued merged range covers it."""
    if any(not _is_empty_cell(value) for value in values):
        return False
    return not any(start <= row_index <= end for start, end in valued_merged_ranges)


def _is_empty_cell(value: object) -> bool:
    return value is None or value == ""


def _blocks_for_sheet(sheet: Worksheet) -> list[BlockSegment]:
    max_row = int(sheet.max_row or 0)
    max_col = int(sheet.max_column or 0)
    if max_row == 0 or max_col == 0:
        return []
    spans = _valued_merged_row_spans(sheet)
    rows = list(sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col))
    return [
        _block_from_run(sheet, start, end, max_col)
        for start, end in _nonempty_runs(rows, spans)
    ]


def _valued_merged_row_spans(sheet: Worksheet) -> tuple[tuple[int, int], ...]:
    spans: list[tuple[int, int]] = []
    for merged in sheet.merged_cells.ranges:
        anchor = sheet.cell(merged.min_row, merged.min_col)
        if not _is_empty_cell(anchor.value):
            spans.append((int(merged.min_row), int(merged.max_row)))
    return tuple(spans)


def _nonempty_runs(
    rows: list[tuple[object, ...]],
    spans: Sequence[tuple[int, int]],
) -> list[tuple[int, int]]:
    runs: list[tuple[int, int]] = []
    start: int | None = None
    for index, row in enumerate(rows, start=1):
        values = tuple(getattr(cell, "value", None) for cell in row)
        if _row_is_empty(index, values, spans):
            if start is not None:
                runs.append((start, index - 1))
                start = None
        elif start is None:
            start = index
    if start is not None:
        runs.append((start, len(rows)))
    return runs


def _block_from_run(
    sheet: Worksheet, start: int, end: int, max_col: int
) -> BlockSegment:
    min_col, used_col = _used_columns(sheet, start, end, max_col)
    label = ""
    header_row: tuple[str, ...] = ()
    header_seen = False
    data_rows: list[tuple[str, ...]] = []
    for row_index in range(start, end + 1):
        texts = tuple(
            _cell_text(sheet, row_index, column)
            for column in range(min_col, used_col + 1)
        )
        if row_index == start:
            label = texts[0] if texts else ""
        if not header_seen:
            if any(texts[1:]):
                header_row = texts
                header_seen = True
            continue
        if _row_has_own_value(sheet, row_index, min_col, used_col):
            data_rows.append(texts)
    cell_range = (
        f"{get_column_letter(min_col)}{start}:{get_column_letter(used_col)}{end}"
    )
    return BlockSegment(
        label=label,
        header_row=header_row,
        rows=tuple(data_rows),
        locator=SheetLocator(sheet=sheet.title, cell_range=cell_range),
    )


def _used_columns(
    sheet: Worksheet, start: int, end: int, max_col: int
) -> tuple[int, int]:
    min_used: int | None = None
    max_used: int | None = None

    def _span(column: int) -> None:
        nonlocal min_used, max_used
        min_used = column if min_used is None else min(min_used, column)
        max_used = column if max_used is None else max(max_used, column)

    for row_index in range(start, end + 1):
        for column in range(1, max_col + 1):
            cell = sheet.cell(row_index, column)
            if isinstance(cell, MergedCell):
                continue
            if not _is_empty_cell(cell.value):
                _span(column)
    for merged in sheet.merged_cells.ranges:
        if int(merged.max_row) < start or int(merged.min_row) > end:
            continue
        anchor = sheet.cell(merged.min_row, merged.min_col)
        if _is_empty_cell(anchor.value):
            continue
        _span(int(merged.min_col))
        _span(int(merged.max_col))
    if min_used is None or max_used is None:
        return 1, 1
    return min_used, max_used


def _row_has_own_value(
    sheet: Worksheet, row_index: int, min_col: int, max_col: int
) -> bool:
    for column in range(min_col, max_col + 1):
        cell = sheet.cell(row_index, column)
        if isinstance(cell, MergedCell):
            continue
        if not _is_empty_cell(cell.value):
            return True
    return False


def _cell_text(sheet: Worksheet, row_index: int, column: int) -> str:
    cell = sheet.cell(row_index, column)
    value = cell.value
    if _is_empty_cell(value):
        return ""
    if isinstance(value, bool):
        return normalise(str(value))
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    if isinstance(value, str):
        return normalise(value)
    return normalise(str(value))

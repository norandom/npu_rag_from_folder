"""Unit tests for the Excel extractor (task 4.4).

Requirement 4.1: a run of entirely empty rows is the block boundary; a row is
empty only if every cell is empty and no valued merged range covers it.
Requirement 4.2: each block carries its label, period header row, and row
labels. Requirement 4.8: a hidden sheet is skipped and recorded by name, never
raised. Requirement 7.3: every block carries a SheetLocator. The committed
workbook is built by generate.py; later tasks read it only.
"""

from __future__ import annotations

import ast
import importlib.util
import zipfile
from collections.abc import Callable, Sequence
from io import BytesIO
from pathlib import Path
from typing import Protocol, cast

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.cell.cell import MergedCell  # type: ignore[import-untyped]

from npu_rag.ingest.config import IngestConfig
from npu_rag.ingest.errors import ExtractionError
from npu_rag.ingest.extract.base import Extractor
from npu_rag.ingest.types import (
    BlockSegment,
    Extracted,
    Omission,
    OmissionCategory,
    SheetLocator,
    SourceFile,
)

REPO_ROOT = Path(__file__).resolve().parents[3]
PACKAGE_ROOT = REPO_ROOT / "src" / "npu_rag" / "ingest"
EXCEL_PATH = PACKAGE_ROOT / "extract" / "excel.py"
FIXTURE_DIR = Path(__file__).resolve().parents[1] / "fixtures" / "excel"
FIXTURE_PATH = FIXTURE_DIR / "models.xlsx"
GENERATE_PATH = FIXTURE_DIR / "generate.py"

LONG_PREFIX = "\\\\?\\"
TRADITIONAL_LIMIT = 260
BUDGET = 256
PNG_MAGIC = b"\x89PNG\r\n\x1a\n"

FORBIDDEN_IMPORTS = (
    "npu_rag.ingest.vision",
    "npu_rag.ingest.state",
    "httpx",
)


class _ExcelFixtureGenerate(Protocol):
    INCOME_SHEET: str
    HIDDEN_SHEET: str
    MERGED_LABEL: str
    MERGED_LABEL_RANGE: str
    FIRST_BLOCK_LABEL: str
    FIRST_BLOCK_RANGE: str
    FIRST_BLOCK_HEADER: tuple[str, ...]
    FIRST_BLOCK_ROWS: tuple[tuple[str, ...], ...]
    SECOND_BLOCK_LABEL: str
    SECOND_BLOCK_RANGE: str
    SECOND_BLOCK_HEADER: tuple[str, ...]
    SECOND_BLOCK_ROWS: tuple[tuple[str, ...], ...]
    CACHED_FORMULA_CELL: str
    CACHED_FORMULA: str
    CACHED_VALUE: int
    UNCACHED_FORMULA_CELL: str
    UNCACHED_FORMULA: str
    CHART_ANCHOR: str
    IMAGE_ANCHOR: str

    def committed_fixture_bytes(self) -> bytes: ...


def load_generate() -> _ExcelFixtureGenerate:
    spec = importlib.util.spec_from_file_location(
        "ingest_excel_fixture_generate", GENERATE_PATH
    )
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return cast(_ExcelFixtureGenerate, module)


GENERATE = load_generate()


def as_windows_long_path(path: Path) -> Path:
    """Discoverer always prefixes ``SourceFile.path`` with ``\\\\?\\`` on Windows."""
    text = str(path if path.is_absolute() else path.resolve())
    if text.startswith("\\\\?\\UNC\\") or text.startswith("\\\\?\\"):
        return Path(text)
    if text.startswith("\\\\"):
        return Path("\\\\?\\UNC\\" + text[2:])
    return Path(LONG_PREFIX + text)


def _without_long_prefix(path: Path | str) -> str:
    text = str(path)
    if text.startswith("\\\\?\\UNC\\"):
        return "\\\\" + text[8:]
    if text.startswith("\\\\?\\"):
        return text[4:]
    return text


def traditional_length(path: Path) -> int:
    return len(_without_long_prefix(path))


def make_config(**overrides: object) -> IngestConfig:
    values: dict[str, object] = {
        "roots": (Path("inputs"),),
        "token_budget": BUDGET,
    }
    values.update(overrides)
    return IngestConfig(**values)  # type: ignore[arg-type]


def plant(root: Path, relative: str, data: bytes) -> SourceFile:
    path = root / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    relative_path = Path(relative)
    author = relative_path.parts[0] if len(relative_path.parts) > 1 else ""
    return SourceFile(
        path=as_windows_long_path(path),
        root=as_windows_long_path(root),
        relative_path=relative_path,
        author=author,
    )


def source_for_fixture() -> SourceFile:
    path = FIXTURE_PATH.resolve()
    return SourceFile(
        path=as_windows_long_path(path),
        root=as_windows_long_path(path.parent),
        relative_path=Path("alice") / "models.xlsx",
        author="alice",
    )


def extract(source: SourceFile, **overrides: object) -> Extracted:
    from npu_rag.ingest.extract.excel import ExcelExtractor

    extractor: Extractor = ExcelExtractor()
    return extractor.extract(source, make_config(**overrides))


def imported_names(path: Path) -> list[str]:
    containing = "npu_rag.ingest.extract"
    parts = containing.split(".")
    names: list[str] = []
    for node in ast.walk(ast.parse(path.read_text(encoding="utf-8"))):
        if isinstance(node, ast.Import):
            names.extend(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom):
            if node.level == 0:
                root = node.module or ""
            else:
                base = ".".join(parts[: len(parts) - node.level + 1])
                root = f"{base}.{node.module}" if node.module else base
            names.append(root)
            names.extend(f"{root}.{alias.name}" for alias in node.names)
    return names


def blocks_of(result: Extracted) -> list[BlockSegment]:
    return [segment for segment in result.segments if isinstance(segment, BlockSegment)]


def _naive_row_is_empty(
    row_index: int,
    values: Sequence[object],
    valued_merged_ranges: Sequence[tuple[int, int]],
) -> bool:
    """All-empty-cells rule: ignores merged coverage. Rows 2-3 of the fixture
    are all-None at the cell level, so this splits the three-row merged label.
    """
    del row_index, valued_merged_ranges
    return all(value is None or value == "" for value in values)


# --------------------------------------------------------------------------
# Fixture non-vacuity
# --------------------------------------------------------------------------


def test_the_generator_reproduces_the_committed_workbook_byte_for_byte() -> None:
    assert GENERATE.committed_fixture_bytes() == FIXTURE_PATH.read_bytes()


def test_the_committed_fixture_still_has_two_blocks_a_three_row_merge_formulas_hidden_sheet_chart_and_image() -> None:
    """Non-vacuity: the on-disk workbook still has every ingredient later tasks consume."""
    raw = FIXTURE_PATH.read_bytes()
    assert raw == GENERATE.committed_fixture_bytes()
    assert raw[:4] == b"PK\x03\x04"

    names = zipfile.ZipFile(BytesIO(raw)).namelist()
    assert any(name.startswith("xl/charts/") for name in names)
    media = [name for name in names if name.startswith("xl/media/")]
    assert len(media) == 1
    image_bytes = zipfile.ZipFile(BytesIO(raw)).read(media[0])
    assert image_bytes.startswith(PNG_MAGIC)

    workbook = load_workbook(BytesIO(raw), data_only=False)
    try:
        assert GENERATE.INCOME_SHEET in workbook.sheetnames
        assert GENERATE.HIDDEN_SHEET in workbook.sheetnames
        income = workbook[GENERATE.INCOME_SHEET]
        notes = workbook[GENERATE.HIDDEN_SHEET]
        assert income.sheet_state == "visible"
        assert notes.sheet_state == "hidden"
        merges = {str(item) for item in income.merged_cells.ranges}
        assert GENERATE.MERGED_LABEL_RANGE in merges
        assert income["A1"].value == GENERATE.MERGED_LABEL
        assert income["A2"].value is None
        assert income["A3"].value is None
        assert isinstance(income["A2"], MergedCell)
        assert isinstance(income["A3"], MergedCell)
        for row in income.iter_rows(min_row=2, max_row=3, max_col=4):
            assert all(cell.value is None for cell in row)
        assert income[GENERATE.CACHED_FORMULA_CELL].value == GENERATE.CACHED_FORMULA
        assert income[GENERATE.UNCACHED_FORMULA_CELL].value == GENERATE.UNCACHED_FORMULA
        assert len(income._charts) == 1
        assert len(income._images) == 1
    finally:
        workbook.close()

    values_only = load_workbook(BytesIO(raw), data_only=True)
    try:
        income = values_only[GENERATE.INCOME_SHEET]
        assert income[GENERATE.CACHED_FORMULA_CELL].value == GENERATE.CACHED_VALUE
        assert income[GENERATE.UNCACHED_FORMULA_CELL].value is None
        assert income[GENERATE.CACHED_FORMULA_CELL].value != income[
            GENERATE.UNCACHED_FORMULA_CELL
        ].value
    finally:
        values_only.close()


def test_a_naive_all_empty_cells_scan_of_the_fixture_splits_the_merged_label() -> None:
    """The three-row merge is stored only in the anchor; all-None rows 2-3
    would be a boundary if merged coverage were ignored.
    """
    workbook = load_workbook(FIXTURE_PATH, data_only=False)
    try:
        sheet = workbook[GENERATE.INCOME_SHEET]
        naive_blocks = _naive_row_runs(sheet)
    finally:
        workbook.close()
    assert naive_blocks != [(1, 6), (9, 12)]
    assert naive_blocks[0] == (1, 1)
    assert naive_blocks[0][1] < 3
    assert len(naive_blocks) == 3


def _naive_row_runs(sheet: object) -> list[tuple[int, int]]:
    max_row = int(getattr(sheet, "max_row"))
    max_col = int(getattr(sheet, "max_column"))
    iter_rows = getattr(sheet, "iter_rows")
    runs: list[tuple[int, int]] = []
    start: int | None = None
    for row_index, row in enumerate(
        iter_rows(min_row=1, max_row=max_row, max_col=max_col), start=1
    ):
        empty = all(cell.value is None or cell.value == "" for cell in row)
        if empty:
            if start is not None:
                runs.append((start, row_index - 1))
                start = None
        elif start is None:
            start = row_index
    if start is not None:
        runs.append((start, max_row))
    return runs


# --------------------------------------------------------------------------
# Requirements 4.1, 4.2, 4.8, 7.3: the fixture observable
# --------------------------------------------------------------------------


def test_the_fixture_yields_exactly_two_blocks_with_the_expected_ranges() -> None:
    result = extract(source_for_fixture())
    blocks = blocks_of(result)
    assert len(blocks) == 2

    first, second = blocks
    assert first.label == GENERATE.FIRST_BLOCK_LABEL
    assert first.header_row == GENERATE.FIRST_BLOCK_HEADER
    assert first.rows == GENERATE.FIRST_BLOCK_ROWS
    assert first.locator == SheetLocator(
        sheet=GENERATE.INCOME_SHEET, cell_range=GENERATE.FIRST_BLOCK_RANGE
    )
    assert first.rows[0][0] == "Product"
    assert first.rows[1][0] == "Services"
    assert first.rows[2][0] == "Total"

    assert second.label == GENERATE.SECOND_BLOCK_LABEL
    assert second.header_row == GENERATE.SECOND_BLOCK_HEADER
    assert second.rows == GENERATE.SECOND_BLOCK_ROWS
    assert second.locator == SheetLocator(
        sheet=GENERATE.INCOME_SHEET, cell_range=GENERATE.SECOND_BLOCK_RANGE
    )
    assert second.rows[0][0] == "COGS"
    assert second.rows[1][0] == "OpEx"
    assert second.rows[2][0] == "Uncached"


def test_every_block_from_the_fixture_carries_a_sheet_locator() -> None:
    result = extract(source_for_fixture())
    blocks = blocks_of(result)
    assert blocks
    for block in blocks:
        assert isinstance(block.locator, SheetLocator)
        assert block.locator.sheet == GENERATE.INCOME_SHEET
        assert ":" in block.locator.cell_range


def test_the_hidden_sheet_is_recorded_as_a_named_omission_and_is_not_raised() -> None:
    source = source_for_fixture()
    result = extract(source)
    assert result.omissions
    hidden = [
        item
        for item in result.omissions
        if item.category is OmissionCategory.HIDDEN_SHEET
    ]
    assert len(hidden) == 1
    omission = hidden[0]
    assert isinstance(omission, Omission)
    assert omission.path == source.path
    assert GENERATE.HIDDEN_SHEET in omission.reason
    assert not any(
        isinstance(segment, BlockSegment)
        and isinstance(segment.locator, SheetLocator)
        and segment.locator.sheet == GENERATE.HIDDEN_SHEET
        for segment in result.segments
    )


def test_a_naive_all_empty_cells_rule_planted_in_place_splits_the_merged_label(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Observable: planting the all-None rule in place of the merged-aware
    rule splits the three-row merged label and no longer yields two blocks.
    """
    from npu_rag.ingest.extract import excel as excel_module

    monkeypatch.setattr(excel_module, "_row_is_empty", _naive_row_is_empty)
    result = extract(source_for_fixture())
    blocks = blocks_of(result)
    assert len(blocks) != 2
    assert len(blocks) == 3
    assert blocks[0].locator == SheetLocator(
        sheet=GENERATE.INCOME_SHEET, cell_range="A1:D1"
    )
    assert blocks[0].locator.cell_range != GENERATE.FIRST_BLOCK_RANGE


def test_excel_extractor_opens_the_windows_long_path_form() -> None:
    source = source_for_fixture()
    assert str(source.path).startswith(LONG_PREFIX)
    result = extract(source)
    assert len(blocks_of(result)) == 2


def test_a_workbook_behind_a_path_longer_than_the_traditional_limit_is_extracted(
    tmp_path: Path,
) -> None:
    root = as_windows_long_path(tmp_path / "archive")
    root.mkdir(parents=True, exist_ok=True)
    directory = root
    name = "models.xlsx"
    while traditional_length(directory / name) <= TRADITIONAL_LIMIT:
        directory = directory / ("subdir" * 4)
        directory.mkdir(parents=True, exist_ok=True)
        if traditional_length(directory) > 400:
            pytest.fail(
                f"cannot construct a path longer than {TRADITIONAL_LIMIT} under {root}"
            )
    target = directory / name
    target.write_bytes(FIXTURE_PATH.read_bytes())
    assert traditional_length(target) > TRADITIONAL_LIMIT
    relative = Path(_without_long_prefix(target)).relative_to(
        Path(_without_long_prefix(root))
    )
    source = SourceFile(
        path=as_windows_long_path(target),
        root=root,
        relative_path=relative,
        author="archive",
    )
    result = extract(source)
    assert len(blocks_of(result)) == 2
    assert blocks_of(result)[0].locator == SheetLocator(
        sheet=GENERATE.INCOME_SHEET, cell_range=GENERATE.FIRST_BLOCK_RANGE
    )


def test_a_corrupt_workbook_raises_extraction_error_with_stage_and_path(
    tmp_path: Path,
) -> None:
    source = plant(tmp_path, "alice/broken.xlsx", b"this is not a workbook")
    with pytest.raises(ExtractionError) as caught:
        extract(source)
    error = caught.value
    assert error.stage == "extraction"
    assert error.path == source.path
    assert str(source.path).startswith(LONG_PREFIX)


def test_an_unreadable_path_raises_extraction_error_with_stage_and_path(
    tmp_path: Path,
) -> None:
    missing = tmp_path / "alice" / "gone.xlsx"
    source = SourceFile(
        path=as_windows_long_path(missing),
        root=as_windows_long_path(tmp_path),
        relative_path=Path("alice") / "gone.xlsx",
        author="alice",
    )
    with pytest.raises(ExtractionError) as caught:
        extract(source)
    error = caught.value
    assert error.stage == "extraction"
    assert error.path == source.path


def test_excel_extractor_does_not_import_vision_state_or_httpx() -> None:
    names = imported_names(EXCEL_PATH)
    for forbidden in FORBIDDEN_IMPORTS:
        assert all(
            name != forbidden and not name.startswith(f"{forbidden}.")
            for name in names
        ), names


def test_row_is_empty_is_the_hook_the_naive_plant_replaces() -> None:
    """The merged-aware rule is a module-level function so the plant is in place."""
    from npu_rag.ingest.extract.excel import _row_is_empty

    hook: Callable[..., bool] = _row_is_empty
    assert hook(2, (None, None, None, None), ((1, 3),)) is False
    assert hook(7, (None, None, None, None), ((1, 3),)) is True
    assert hook(4, (None, "100", None, None), ()) is False

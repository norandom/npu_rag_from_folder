"""Rebuild the committed Excel fixture workbook (task 4.4).

openpyxl timestamps zip members and never writes a formula's cached value, so
this script builds the workbook, injects one cached ``<v>``, and re-packs the
archive with a fixed date. Re-running this file must reproduce ``models.xlsx``
byte for byte.
"""

from __future__ import annotations

import datetime
import re
import zipfile
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.chart import BarChart, Reference  # type: ignore[import-untyped]
from openpyxl.drawing.image import Image as XLImage  # type: ignore[import-untyped]
from openpyxl.packaging.core import DocumentProperties  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]

__all__ = [
    "CACHED_FORMULA",
    "CACHED_FORMULA_CELL",
    "CACHED_VALUE",
    "CHART_ANCHOR",
    "FIRST_BLOCK_HEADER",
    "FIRST_BLOCK_LABEL",
    "FIRST_BLOCK_RANGE",
    "FIRST_BLOCK_ROWS",
    "HIDDEN_SHEET",
    "IMAGE_ANCHOR",
    "INCOME_SHEET",
    "MERGED_LABEL",
    "MERGED_LABEL_RANGE",
    "SECOND_BLOCK_HEADER",
    "SECOND_BLOCK_LABEL",
    "SECOND_BLOCK_RANGE",
    "SECOND_BLOCK_ROWS",
    "UNCACHED_FORMULA",
    "UNCACHED_FORMULA_CELL",
    "build_workbook_bytes",
    "committed_fixture_bytes",
]

_FIXTURE_NAME = "models.xlsx"
_ZIP_DATE = (2026, 1, 1, 0, 0, 0)
_CREATED = datetime.datetime(2026, 1, 1, 0, 0, 0)

INCOME_SHEET = "Income"
HIDDEN_SHEET = "Notes"
MERGED_LABEL = "Revenue"
MERGED_LABEL_RANGE = "A1:A3"
FIRST_BLOCK_LABEL = "Revenue"
FIRST_BLOCK_RANGE = "A1:D6"
FIRST_BLOCK_HEADER = ("Revenue", "FY24", "FY25", "FY26")
FIRST_BLOCK_ROWS = (
    ("Product", "100", "110", "121"),
    ("Services", "50", "55", "61"),
    ("Total", "=B4+B5", "=C4+C5", "=D4+D5"),
)
SECOND_BLOCK_LABEL = "Expenses"
SECOND_BLOCK_RANGE = "A9:D12"
SECOND_BLOCK_HEADER = ("Expenses", "Q1", "Q2", "Q3")
SECOND_BLOCK_ROWS = (
    ("COGS", "30", "32", "34"),
    ("OpEx", "20", "21", "22"),
    ("Uncached", "=B10+B11", "", ""),
)
CACHED_FORMULA_CELL = "B6"
CACHED_FORMULA = "=B4+B5"
CACHED_VALUE = 150
UNCACHED_FORMULA_CELL = "B12"
UNCACHED_FORMULA = "=B10+B11"
CHART_ANCHOR = "F1"
IMAGE_ANCHOR = "F15"

# 8x8 RGB PNG, frozen so the media part does not depend on Pillow's encoder.
_PNG_BYTES = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000080000000808020000004b6d29dc"
    "0000001449444154789c63640858c0800d3061151db41200c0080100b7d11c10"
    "0000000049454e44ae426082"
)

_CACHED_FORMULA_XML = re.compile(
    rb'(<c r="B6"[^>]*>\s*<f>B4\+B5</f>)\s*<v\s*/>\s*(</c>)'
)
_MODIFIED_XML = re.compile(
    rb"<dcterms:modified xsi:type=\"dcterms:W3CDTF\">[^<]+</dcterms:modified>"
)
_FIXED_MODIFIED = (
    b'<dcterms:modified xsi:type="dcterms:W3CDTF">2026-01-01T00:00:00Z'
    b"</dcterms:modified>"
)


def build_workbook_bytes() -> bytes:
    """Build the fixture workbook bytes, including the cached formula value."""
    workbook = Workbook()
    workbook.properties = DocumentProperties(
        creator="npu_rag",
        lastModifiedBy="npu_rag",
        created=_CREATED,
        modified=_CREATED,
        title="models",
    )
    income = workbook.active
    assert income is not None
    income.title = INCOME_SHEET
    _fill_income(income)
    notes = workbook.create_sheet(HIDDEN_SHEET)
    notes.sheet_state = "hidden"
    notes["A1"] = "Hidden notes"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return _canonical_zip(_inject_cached_value(buffer.getvalue()))


def committed_fixture_bytes() -> bytes:
    return build_workbook_bytes()


def _fill_income(sheet: Worksheet) -> None:
    sheet["A1"] = MERGED_LABEL
    sheet["B1"] = "FY24"
    sheet["C1"] = "FY25"
    sheet["D1"] = "FY26"
    sheet.merge_cells(MERGED_LABEL_RANGE)

    sheet["A4"] = "Product"
    sheet["B4"] = 100
    sheet["C4"] = 110
    sheet["D4"] = 121
    sheet["A5"] = "Services"
    sheet["B5"] = 50
    sheet["C5"] = 55
    sheet["D5"] = 61
    sheet["A6"] = "Total"
    sheet["B6"] = CACHED_FORMULA
    sheet["C6"] = "=C4+C5"
    sheet["D6"] = "=D4+D5"

    sheet["A9"] = "Expenses"
    sheet["B9"] = "Q1"
    sheet["C9"] = "Q2"
    sheet["D9"] = "Q3"
    sheet["A10"] = "COGS"
    sheet["B10"] = 30
    sheet["C10"] = 32
    sheet["D10"] = 34
    sheet["A11"] = "OpEx"
    sheet["B11"] = 20
    sheet["C11"] = 21
    sheet["D11"] = 22
    sheet["A12"] = "Uncached"
    sheet["B12"] = UNCACHED_FORMULA

    chart = BarChart()
    chart.title = "Revenue"
    data = Reference(sheet, min_col=2, min_row=4, max_col=2, max_row=5)
    cats = Reference(sheet, min_col=1, min_row=4, max_row=5)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(cats)
    sheet.add_chart(chart, CHART_ANCHOR)

    image = XLImage(BytesIO(_PNG_BYTES))
    sheet.add_image(image, IMAGE_ANCHOR)


def _inject_cached_value(raw: bytes) -> bytes:
    source = zipfile.ZipFile(BytesIO(raw), "r")
    out = BytesIO()
    with zipfile.ZipFile(out, "w") as dest:
        patched = False
        pinned_modified = False
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename.startswith("xl/worksheets/sheet") and info.filename.endswith(
                ".xml"
            ):
                updated, count = _CACHED_FORMULA_XML.subn(
                    rb"\1<v>" + str(CACHED_VALUE).encode("ascii") + rb"</v>\2",
                    payload,
                    count=1,
                )
                if count:
                    payload = updated
                    patched = True
            if info.filename == "docProps/core.xml":
                payload, count = _MODIFIED_XML.subn(_FIXED_MODIFIED, payload, count=1)
                if count:
                    pinned_modified = True
            dest.writestr(info, payload)
    if not patched:
        raise RuntimeError("cached formula cell B6 was not present in sheet XML")
    if not pinned_modified:
        raise RuntimeError("core.xml modified timestamp was not present")
    return out.getvalue()


def _canonical_zip(raw: bytes) -> bytes:
    source = zipfile.ZipFile(BytesIO(raw), "r")
    out = BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as dest:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(filename=name, date_time=_ZIP_DATE)
            info.compress_type = zipfile.ZIP_DEFLATED
            dest.writestr(info, source.read(name))
    return out.getvalue()


def main() -> None:
    target = Path(__file__).with_name(_FIXTURE_NAME)
    target.write_bytes(committed_fixture_bytes())


if __name__ == "__main__":
    main()
